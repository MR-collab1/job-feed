"""The merged Excel export: nothing dropped, NA everywhere data is missing."""

from __future__ import annotations

import datetime as dt

import openpyxl
import pytest

from alcohol_audit.export import (
    NA,
    build_merged_rows,
    is_empty_value,
    read_sheet,
    write_merged_workbook,
)


def rows_from(make_record, *overrides):
    from alcohol_audit.etl import COLUMNS

    names = [name for _, name in COLUMNS]
    return [[make_record(**o)[name] for name in names] for o in overrides]


@pytest.mark.parametrize("value", ["", None, "-", "--", "N/A", "n/a", "NA", "?", "  "])
def test_empty_tokens_become_na(value):
    assert is_empty_value(value) is True


@pytest.mark.parametrize(
    "value", ["No", "Yes", "case Notes missing", "Not documented", "Unclear handwriting", "0"]
)
def test_meaningful_text_is_not_treated_as_empty(value):
    assert is_empty_value(value) is False


def test_every_source_row_survives_the_merge(make_record):
    audit = rows_from(make_record,
                      {"patient_id": "A", "arrival": "01/12/2025 10:00"},
                      {"patient_id": "B", "arrival": "02/12/2025 10:00"})
    complaint = rows_from(make_record, {"patient_id": "A", "arrival": "01/12/2025 10:00",
                                        "outcome_raw": "Admitted"})
    merged, counts = build_merged_rows(audit, complaint)

    assert counts["output_rows"] == 2
    assert counts["matched"] == 1
    assert counts["audit_only"] == 1
    assert counts["complaint_only"] == 0


def test_a_row_only_in_sheet3_is_still_carried_through(make_record):
    complaint = rows_from(make_record, {"patient_id": "Z", "arrival": "09/09/2026 10:00"})
    merged, counts = build_merged_rows([], complaint)

    assert counts["output_rows"] == 1
    assert counts["complaint_only"] == 1
    assert merged[0]["in_audit"] == "No"


def test_section_heading_row_is_kept_and_flagged(make_record):
    audit = rows_from(make_record, {"patient_id": "A", "arrival": "01/12/2025 10:00"})
    audit.append(["1st November 2025 to 27th December 2025"] + [None] * 18)
    merged, counts = build_merged_rows(audit, [])

    assert counts["output_rows"] == 2
    assert counts["section_headings"] == 1
    heading = [r for r in merged if r["record_type"] != "Attendance"]
    assert len(heading) == 1
    assert heading[0]["clinical_source"] == "Not applicable"


def test_sheet3_fills_blanks_without_overwriting_existing_values(make_record):
    audit = rows_from(make_record, {"patient_id": "A", "arrival": "01/12/2025 10:00",
                                    "outcome_raw": "Admitted", "ciwa_score_raw": "-"})
    complaint = rows_from(make_record, {"patient_id": "A", "arrival": "01/12/2025 10:00",
                                        "outcome_raw": "N/A", "ciwa_score_raw": "17"})
    merged, counts = build_merged_rows(audit, complaint)

    values = merged[0]["values"]
    assert values[8] == "17"        # blank in the audit sheet, filled from Sheet3
    assert values[18] == "Admitted"  # already present, and Sheet3 had nothing
    assert counts["cells_filled_from_complaint"] == 1


def test_written_workbook_has_no_empty_cells(tmp_path, make_record):
    audit = rows_from(make_record,
                      {"patient_id": "A", "arrival": "01/12/2025 10:00"},
                      {"patient_id": "B", "arrival": "02/12/2025 10:00", "age": None})
    path, counts = write_merged_workbook(audit, [], tmp_path / "merged.xlsx")

    sheet = openpyxl.load_workbook(path)["Merged Audit"]
    assert sheet.max_row == counts["output_rows"] + 1
    for row in range(2, sheet.max_row + 1):
        for column in range(1, sheet.max_column + 1):
            assert sheet.cell(row=row, column=column).value not in (None, "")


def test_dates_are_written_as_real_datetimes(tmp_path, make_record):
    audit = rows_from(make_record, {"patient_id": "A", "arrival": "01/12/2025 10:00",
                                    "departure": "01/12/2025 14:30"})
    path, _ = write_merged_workbook(audit, [], tmp_path / "merged.xlsx")
    sheet = openpyxl.load_workbook(path)["Merged Audit"]

    assert sheet.cell(row=2, column=3).value == dt.datetime(2025, 12, 1, 10, 0)
    assert sheet.cell(row=2, column=4).value == dt.datetime(2025, 12, 1, 14, 30)


def test_missing_date_is_na_not_a_broken_date(tmp_path, make_record):
    audit = rows_from(make_record, {"patient_id": "A", "arrival": None, "departure": None})
    path, _ = write_merged_workbook(audit, [], tmp_path / "merged.xlsx")
    sheet = openpyxl.load_workbook(path)["Merged Audit"]

    assert sheet.cell(row=2, column=3).value == NA


def test_notes_sheet_is_present(tmp_path, make_record):
    audit = rows_from(make_record, {"patient_id": "A", "arrival": "01/12/2025 10:00"})
    path, _ = write_merged_workbook(audit, [], tmp_path / "merged.xlsx")
    assert openpyxl.load_workbook(path).sheetnames == ["Merged Audit", "Merge Notes"]


def test_real_export_keeps_every_source_row(tmp_path, real_dataset):
    """Against the real workbooks: output rows equal the source rows carrying content."""
    from alcohol_audit.etl import AUDIT_SHEET, COMPLAINT_SHEET, _resolve, default_data_dir

    directory = default_data_dir()
    audit_path = _resolve(directory, "ALCOHOL_AUDIT_CIWA_FILE", "*Audit_CIWA*.xlsx")
    complaint_path = _resolve(
        directory, "ALCOHOL_AUDIT_COMPLAINT_FILE", "*PresentingComplaint*.xlsx"
    )
    headers, audit_rows = read_sheet(audit_path, AUDIT_SHEET)
    _, complaint_rows = read_sheet(complaint_path, COMPLAINT_SHEET)

    path, counts = write_merged_workbook(
        audit_rows, complaint_rows, tmp_path / "merged.xlsx", headers=headers
    )

    # Every attendance in the analytics dataset appears, plus the heading row.
    assert counts["output_rows"] == len(real_dataset) + counts["section_headings"]
    assert counts["complaint_only"] == 0
    assert counts["matched"] == counts["complaint_rows"]

    sheet = openpyxl.load_workbook(path)["Merged Audit"]
    assert sheet.cell(row=1, column=11).value.startswith("Which treatment was given?")
    for row in range(2, sheet.max_row + 1):
        for column in range(1, sheet.max_column + 1):
            assert sheet.cell(row=row, column=column).value not in (None, "")
