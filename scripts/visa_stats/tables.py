"""Turn a downloaded CSV/XLSX resource into a list of row dicts.

Home Affairs publishes most of its statistics as spreadsheets that were built
for humans: a title row, a blank row, sometimes a footnote block, and only then
the actual header. Everything here exists to find the real header row and hand
back plain records.
"""

from __future__ import annotations

import csv
import io
import logging
import re
from dataclasses import dataclass, field

log = logging.getLogger(__name__)

# How far into a sheet we are willing to look for the header row.
MAX_HEADER_SCAN = 25

# Cell values that mean "no data" in these publications.
NULL_TOKENS = {"", "-", "--", "n/a", "na", "n.a.", "np", "n.p.", "..", "nil", "none"}


class TableError(RuntimeError):
    """Raised when a resource could not be parsed into a table."""


@dataclass
class Table:
    """A parsed tabular resource."""

    headers: list[str]
    rows: list[dict[str, str]]
    sheet: str | None = None
    notes: list[str] = field(default_factory=list)

    def __len__(self) -> int:
        return len(self.rows)


def normalise_header(value: object) -> str:
    """Collapse a header cell to a stable lowercase key.

    ``"Country of\\nCitizenship "`` and ``"country of citizenship"`` both become
    ``"country of citizenship"``, so the column matchers only have to describe
    one spelling.
    """
    text = "" if value is None else str(value)
    text = text.replace(" ", " ")
    text = re.sub(r"\[\s*\d+\s*\]|\(\s*[a-z]\s*\)$", " ", text, flags=re.IGNORECASE)
    text = re.sub(r"[^0-9a-zA-Z%/&+.-]+", " ", text)
    return re.sub(r"\s+", " ", text).strip().lower()


def clean_cell(value: object) -> str:
    """Trim a data cell and normalise the various 'no value' spellings to ''."""
    if value is None:
        return ""
    text = str(value).replace(" ", " ").strip()
    return "" if text.lower() in NULL_TOKENS else text


def parse_number(value: object) -> float | None:
    """Parse ``"12,345"``, ``"1 234"``, ``"45.6%"``, ``"<5"`` into a float.

    ``"<5"`` appears in these datasets where a count is suppressed for privacy;
    it is treated as unknown rather than silently read as 5.
    """
    text = clean_cell(value)
    if not text or text.startswith("<"):
        return None
    text = text.replace(",", "").replace(" ", "").replace("%", "")
    if text.startswith("(") and text.endswith(")"):  # accounting negative
        text = "-" + text[1:-1]
    try:
        return float(text)
    except ValueError:
        return None


def read_csv(body: bytes) -> Table:
    """Parse CSV/TSV bytes, skipping any preamble above the header row."""
    text = body.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","

    grid = [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]
    return _table_from_grid(grid, sheet=None)


def read_xlsx(body: bytes, *, sheet_pattern: str | None = None) -> Table:
    """Parse XLSX bytes, choosing a sheet and skipping any preamble rows."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - CI installs openpyxl
        raise TableError("openpyxl is required to read .xlsx resources") from exc

    workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    try:
        sheet = _choose_sheet(workbook, sheet_pattern)
        grid = [
            ["" if cell is None else cell for cell in row]
            for row in sheet.iter_rows(max_row=5000, values_only=True)
        ]
    finally:
        workbook.close()

    return _table_from_grid(grid, sheet=sheet.title)


def read_resource(body: bytes, *, fmt: str, sheet_pattern: str | None = None) -> Table:
    """Dispatch to the right parser for a resource format string."""
    kind = (fmt or "").strip().lower().lstrip(".")
    if kind in {"xlsx", "xlsm", "xls"}:
        return read_xlsx(body, sheet_pattern=sheet_pattern)
    if kind in {"csv", "tsv", "txt", ""}:
        return read_csv(body)
    raise TableError(f"unsupported resource format: {fmt!r}")


def table_from_records(records: list[dict], fields: list[str] | None = None) -> Table:
    """Build a Table from CKAN datastore records (already key/value shaped)."""
    if fields is None:
        seen: dict[str, None] = {}
        for record in records:
            for key in record:
                seen.setdefault(key, None)
        fields = list(seen)

    headers = [normalise_header(f) for f in fields]
    rows = [
        {
            normalise_header(f): clean_cell(record.get(f))
            for f in fields
            if normalise_header(f)
        }
        for record in records
    ]
    return Table(headers=[h for h in headers if h], rows=rows)


def _choose_sheet(workbook, sheet_pattern: str | None):
    """Pick the sheet matching ``sheet_pattern``, else the first data-bearing one."""
    if sheet_pattern:
        matcher = re.compile(sheet_pattern, re.IGNORECASE)
        for name in workbook.sheetnames:
            if matcher.search(name):
                return workbook[name]
        log.warning("no sheet matched %r; falling back to first sheet", sheet_pattern)

    for name in workbook.sheetnames:
        # "Notes", "Contents" and "Explanatory notes" tabs lead these workbooks.
        if not re.search(r"note|content|cover|about|glossar|disclaimer", name, re.IGNORECASE):
            return workbook[name]
    return workbook[workbook.sheetnames[0]]


def _table_from_grid(grid: list[list], *, sheet: str | None) -> Table:
    """Locate the header row in a raw grid and build records below it."""
    if not grid:
        raise TableError("resource contained no rows")

    header_index = _find_header_row(grid)
    if header_index is None:
        raise TableError("could not locate a header row")

    headers = _dedupe([normalise_header(cell) for cell in grid[header_index]])
    notes = [
        str(row[0]).strip()
        for row in grid[:header_index]
        if row and str(row[0]).strip()
    ]

    rows: list[dict[str, str]] = []
    for raw in grid[header_index + 1:]:
        record = {
            header: clean_cell(raw[i]) if i < len(raw) else ""
            for i, header in enumerate(headers)
            if header
        }
        if any(record.values()):
            rows.append(record)

    return Table(
        headers=[h for h in headers if h],
        rows=rows,
        sheet=sheet,
        notes=notes[:5],
    )


def _find_header_row(grid: list[list]) -> int | None:
    """Score candidate rows and return the index of the most header-like one.

    A header row has several distinct, short, non-numeric cells and is followed
    by a row that actually carries data.
    """
    best_index: int | None = None
    best_score = 0.0

    for index, row in enumerate(grid[:MAX_HEADER_SCAN]):
        labels = [normalise_header(cell) for cell in row]
        filled = [label for label in labels if label]
        if len(filled) < 2 or len(set(filled)) < 2:
            continue
        # Reject rows that are mostly numbers - that is data, not a header.
        numeric = sum(1 for label in filled if parse_number(label) is not None)
        if numeric > len(filled) / 2:
            continue
        if not _has_data_below(grid, index):
            continue

        score = len(set(filled)) + (1.0 if len(filled) == len(labels) else 0.0)
        if score > best_score:
            best_score, best_index = score, index

    return best_index


def _has_data_below(grid: list[list], index: int) -> bool:
    for row in grid[index + 1: index + 4]:
        if any(clean_cell(cell) for cell in row):
            return True
    return False


def _dedupe(headers: list[str]) -> list[str]:
    """Make repeated header names unique so records do not lose columns."""
    seen: dict[str, int] = {}
    output: list[str] = []
    for header in headers:
        if not header:
            output.append("")
            continue
        if header in seen:
            seen[header] += 1
            output.append(f"{header} {seen[header]}")
        else:
            seen[header] = 1
            output.append(header)
    return output
